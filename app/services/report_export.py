"""Export AI reports to PDF, DOCX, TXT and MD."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date

from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

from app.services.assistant.format_reply import sanitize_llm_plain_text

REPORT_TYPE_LABELS = {
    "clinico": "Relatório Clínico",
    "escolar": "Relatório Escolar",
    "pais": "Relatório para Pais",
    "evolutivo": "Relatório Evolutivo",
    "consolidado": "Laudo Consolidado",
}

_UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9._-]+")


def sanitize_filename_component(value: str) -> str:
    """Keep only header-safe characters for Content-Disposition filenames."""
    cleaned = _UNSAFE_FILENAME_CHARS.sub("-", (value or "").strip()).strip("-.")
    return cleaned or "documento"


_HEADING2 = re.compile(r"^##\s+(.+)$")
_HEADING3 = re.compile(r"^###\s+(.+)$")
_BULLET = re.compile(r"^[-*]\s+(.+)$")
_TABLE_SEPARATOR = re.compile(r"^\|[\s\-:|]+\|?\s*$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


@dataclass(frozen=True)
class DocumentIdentity:
    """Professional identity carried by delivered/exported documents."""

    professional_name: str = ""
    council: str = ""
    issued_at: date | None = None
    logo_bytes: bytes | None = None
    signature_bytes: bytes | None = None

    @property
    def has_content(self) -> bool:
        return bool(self.professional_name or self.council or self.logo_bytes or self.signature_bytes)


def _identity_headline(identity: DocumentIdentity | None) -> str:
    if identity is None:
        return ""
    name = identity.professional_name.strip()
    council = identity.council.strip()
    if name and council:
        return f"{name} — {council}"
    return name or council


def _issued_label(identity: DocumentIdentity | None, fallback: date) -> str:
    issued = identity.issued_at if identity and identity.issued_at else fallback
    return issued.isoformat()


def _format_table_row(line: str) -> str:
    trimmed = line.strip()
    if "|" not in trimmed:
        return line
    inner = trimmed.strip("|")
    cells = [cell.strip() for cell in inner.split("|") if cell.strip()]
    return " · ".join(cells) if cells else ""


def _apply_bold_reportlab(text: str) -> str:
    escaped = text.replace("&", "&amp;").replace("<", "&lt;")
    return _BOLD.sub(r"<b>\1</b>", escaped)


def _iter_markdown_lines(content: str):
    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            yield ("blank", "")
            continue
        if _TABLE_SEPARATOR.match(line):
            continue
        if line.strip().startswith("|"):
            yield ("text", _format_table_row(line))
            continue
        h2 = _HEADING2.match(line)
        if h2:
            yield ("h2", h2.group(1).strip())
            continue
        h3 = _HEADING3.match(line)
        if h3:
            yield ("h3", h3.group(1).strip())
            continue
        bullet = _BULLET.match(line)
        if bullet:
            yield ("bullet", bullet.group(1).strip())
            continue
        yield ("text", line.strip())


def _add_bold_runs(paragraph, text: str) -> None:
    parts = _BOLD.split(text)
    for index, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if index % 2 == 1:
            run.bold = True


def _header_lines(
    report_type: str,
    patient_name: str,
    report_date: date,
    identity: DocumentIdentity | None = None,
) -> list[str]:
    type_label = REPORT_TYPE_LABELS.get(report_type, report_type)
    lines = [
        "KorusFono",
        type_label,
        f"Paciente: {patient_name}",
        f"Data: {report_date.isoformat()}",
    ]
    headline = _identity_headline(identity)
    if headline:
        lines.append(f"Profissional: {headline}")
    lines.append("")
    return lines


def _footer_text_lines(
    report_date: date, identity: DocumentIdentity | None = None
) -> list[str]:
    headline = _identity_headline(identity)
    if not headline:
        return []
    return [
        "",
        "—",
        headline,
        f"Emitido em {_issued_label(identity, report_date)}",
        "",
    ]


def export_txt(
    report_type: str,
    patient_name: str,
    report_date: date,
    content: str,
    identity: DocumentIdentity | None = None,
) -> bytes:
    lines = _header_lines(report_type, patient_name, report_date, identity)
    lines.append(sanitize_llm_plain_text(content))
    lines.extend(_footer_text_lines(report_date, identity))
    return "\n".join(lines).encode("utf-8")


def export_md(
    report_type: str,
    patient_name: str,
    report_date: date,
    content: str,
    identity: DocumentIdentity | None = None,
) -> bytes:
    type_label = REPORT_TYPE_LABELS.get(report_type, report_type)
    headline = _identity_headline(identity)
    professional_line = f"**Profissional:** {headline}  \n" if headline else ""
    body = (
        f"# {type_label}\n\n"
        f"**Paciente:** {patient_name}  \n"
        f"**Data:** {report_date.isoformat()}  \n"
        f"{professional_line}"
        f"\n---\n\n"
        f"{content}\n"
    )
    if headline:
        body += (
            f"\n---\n\n"
            f"{headline}  \n"
            f"Emitido em {_issued_label(identity, report_date)}\n"
        )
    return body.encode("utf-8")


def _docx_add_image(doc, data: bytes | None, *, width_inches: float) -> None:
    if not data:
        return
    try:
        doc.add_picture(io.BytesIO(data), width=Inches(width_inches))
    except Exception:  # noqa: BLE001 - a broken image must not break the document
        pass


def export_docx(
    report_type: str,
    patient_name: str,
    report_date: date,
    content: str,
    identity: DocumentIdentity | None = None,
) -> bytes:
    doc = Document()
    _docx_add_image(doc, identity.logo_bytes if identity else None, width_inches=1.1)
    doc.add_heading("KorusFono", level=1)
    type_label = REPORT_TYPE_LABELS.get(report_type, report_type)
    doc.add_heading(type_label, level=2)
    doc.add_paragraph(f"Paciente: {patient_name}")
    doc.add_paragraph(f"Data: {report_date.isoformat()}")
    headline = _identity_headline(identity)
    if headline:
        doc.add_paragraph(f"Profissional: {headline}")
    doc.add_paragraph("")

    for kind, value in _iter_markdown_lines(content):
        if kind == "blank":
            doc.add_paragraph("")
        elif kind == "h2":
            doc.add_heading(value, level=2)
        elif kind == "h3":
            doc.add_heading(value, level=3)
        elif kind == "bullet":
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_bold_runs(paragraph, value)
        else:
            paragraph = doc.add_paragraph()
            _add_bold_runs(paragraph, value)

    if headline:
        doc.add_paragraph("")
        _docx_add_image(doc, identity.signature_bytes if identity else None, width_inches=1.6)
        doc.add_paragraph(headline)
        doc.add_paragraph(f"Emitido em {_issued_label(identity, report_date)}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pdf_image_flowable(data: bytes | None, *, width: float) -> Image | None:
    if not data:
        return None
    try:
        reader = ImageReader(io.BytesIO(data))
        original_width, original_height = reader.getSize()
        if not original_width or not original_height:
            return None
        return Image(io.BytesIO(data), width=width, height=width * original_height / original_width)
    except Exception:  # noqa: BLE001 - skip images that reportlab cannot decode
        return None


def export_pdf(
    report_type: str,
    patient_name: str,
    report_date: date,
    content: str,
    identity: DocumentIdentity | None = None,
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    type_label = REPORT_TYPE_LABELS.get(report_type, report_type)
    story = []
    logo = _pdf_image_flowable(
        identity.logo_bytes if identity else None, width=110
    )
    if logo is not None:
        story.append(logo)
        story.append(Spacer(1, 8))
    story.extend(
        [
            Paragraph("KorusFono", styles["Title"]),
            Paragraph(type_label, styles["Heading2"]),
            Paragraph(f"Paciente: {patient_name}", styles["Normal"]),
            Paragraph(f"Data: {report_date.isoformat()}", styles["Normal"]),
        ]
    )
    headline = _identity_headline(identity)
    if headline:
        story.append(Paragraph(f"Profissional: {headline}", styles["Normal"]))
    story.append(Spacer(1, 12))
    for kind, value in _iter_markdown_lines(content):
        if kind == "blank":
            story.append(Spacer(1, 6))
        elif kind == "h2":
            story.append(Paragraph(_apply_bold_reportlab(value), styles["Heading2"]))
        elif kind == "h3":
            story.append(Paragraph(_apply_bold_reportlab(value), styles["Heading3"]))
        elif kind == "bullet":
            story.append(Paragraph(f"• {_apply_bold_reportlab(value)}", styles["Normal"]))
        else:
            story.append(Paragraph(_apply_bold_reportlab(value), styles["Normal"]))
    if headline:
        signature = _pdf_image_flowable(
            identity.signature_bytes if identity else None, width=140
        )
        story.append(Spacer(1, 24))
        if signature is not None:
            story.append(signature)
            story.append(Spacer(1, 4))
        story.append(Paragraph(headline, styles["Normal"]))
        story.append(
            Paragraph(
                f"Emitido em {_issued_label(identity, report_date)}",
                styles["Normal"],
            )
        )
    doc.build(story)
    return buffer.getvalue()


def export_report(
    format: str,
    report_type: str,
    patient_name: str,
    report_date: date,
    content: str,
    identity: DocumentIdentity | None = None,
) -> tuple[bytes, str, str]:
    """Return (bytes, media_type, filename_suffix)."""
    if format in ("txt", "md"):
        exporter = export_txt if format == "txt" else export_md
        media = "text/plain; charset=utf-8" if format == "txt" else "text/markdown; charset=utf-8"
        return exporter(report_type, patient_name, report_date, content, identity), media, format
    if format == "docx":
        return (
            export_docx(report_type, patient_name, report_date, content, identity),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    if format == "pdf":
        return (
            export_pdf(report_type, patient_name, report_date, content, identity),
            "application/pdf",
            "pdf",
        )
    raise ValueError(f"Formato não suportado: {format}")
